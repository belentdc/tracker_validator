#!/usr/bin/env python3
"""Build a full, machine- and human-readable NDC Transport Tracker taxonomy
from the source workbook NDC-taxonomy.xlsx.

Outputs:
  - ndc_taxonomy.json   (canonical, machine-readable)
  - ndc_taxonomy.csv    (flat, spreadsheet-friendly)
  - TAXONOMY.md         (human-readable README with Mermaid diagrams)

Definition policy: the original definition is preserved verbatim. Where the
original contains a typo or unclear phrasing, a suggested clarified version is
added in the field 'definition_suggested' so the maintainer can choose. The
original is never altered.
"""
import json
import csv
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent          # taxonomy/build/
OUT = HERE.parent                                # taxonomy/
SRC = HERE / "NDC-taxonomy.xlsx"

wb = load_workbook(SRC, read_only=True, data_only=True)


def rows(sheet):
    ws = wb[sheet]
    for r in ws.iter_rows(values_only=True):
        yield [(str(c).strip() if c is not None else "") for c in r]


# ---------------------------------------------------------------------------
# 1. Build a source lookup (parameter code -> definition source) from
#    Def_Indicators, which is the only sheet carrying the Source column.
# ---------------------------------------------------------------------------
source_lookup = {}
for r in rows("Def_Indicators"):
    # columns: Category | Purpose | Measure | Parameter | Definition | Source
    if len(r) >= 6 and r[3] and r[5] and r[3] != "Parameter":
        source_lookup[r[3]] = r[5]


# ---------------------------------------------------------------------------
# 2. Parse Def_Targets, the canonical machine sheet. Layout (col B onward):
#    Type of indicator | Indicator | Value name | Value identifier |
#    Related (category/target area) | Related (purpose/GHG) | Definition
#    Data begins after a header row per block. We read every data row that has
#    a Value name + Definition, tagging by the 'Type of indicator' block.
# ---------------------------------------------------------------------------
records = []
for r in rows("Def_Targets"):
    # shift: col A is blank, real data starts at index 1
    if len(r) < 8:
        continue
    type_ind, indicator, value_name, value_id, rel1, rel2, definition = (
        r[1], r[2], r[3], r[4], r[5], r[6], r[7]
    )
    if not value_name or value_name in ("Value name", "NA"):
        continue
    if type_ind not in ("Targets", "Mitigation", "Adaptation", "Benefits"):
        continue
    rec = {
        "domain": type_ind,
        "dimension": indicator,
        "label": value_name,
        "code": value_id if value_id and value_id != "NA" else None,
        "parent_1": rel1 if rel1 and rel1 != "NA" else None,
        "parent_2": rel2 if rel2 and rel2 != "NA" else None,
        "definition": definition or None,
    }
    rec["source"] = source_lookup.get(rec["code"]) if rec["code"] else None
    records.append(rec)


# ---------------------------------------------------------------------------
# 3. Clarity suggestions. Original is preserved; suggestion is offered only
#    where the source text has a clear typo or garbled phrasing. Keyed by code
#    (or label when code is absent). Maintainer decides whether to adopt.
# ---------------------------------------------------------------------------
suggestions = {
    "A_Complan": None,  # source typo 'VTI (215)' flagged separately
    "I_Education": "This parameter is for general educational activities and behavioural change related to transport (e.g. the environmental impacts of private vehicle use, the benefits of electric vehicles, etc.).",
    "I_Ecodriving": "Ecodriving refers to educational measures that encourage more efficient driving practices, which can reduce fuel consumption. Captured under this parameter.",
    "A_Caraccess": "This parameter refers to measures that restrict the physical access of certain vehicle types to certain places (e.g. city centres); Low Emission Zones are one example. By restricting access based on criteria such as propulsion technology, Euro standard or age, these measures can deliver differentiated impacts (GHG reduction, lower air pollution, congestion prevention) and can encourage low-carbon modes and a shift to cleaner technologies.",
    "I_VehicleRestrictions": "This parameter encompasses restrictions on vehicle ownership or purchase, including import bans on older vehicles or sale restrictions on particularly polluting vehicles.",
    "A_SUMP": "A sustainable urban mobility plan (SUMP) is a strategic plan designed to satisfy the mobility needs of people and businesses in cities and their surroundings for a better quality of life. It builds on existing planning practices and gives due consideration to integration, participation and evaluation principles. Captured here when a document refers to a SUMP or an integrated urban mobility approach.",
    "S_Activemobility": "General measures that refer to walking and cycling are included here.",
    "S_Walking": "Any action that specifically mentions improving walking is included here.",
    "S_Cycling": "Any action that specifically mentions improving cycling is included here.",
    "A_LATM": "This parameter looks at management, infrastructure and technological approaches with the goal of improving traffic flow.",
    "I_Efficiencystd": "This parameter captures emission standards regulating air-pollutant exhaust emissions (such as NOx), e.g. the EURO standards Euro 1-6 (not referring to CO2 standards).",
    "I_Biofuel": "Conventional diesel and gasoline can be blended with less carbon-intensive fuels. Many national governments set blending mandates (e.g. 10% or 20% biofuel). The most common biofuel is ethanol. General biofuel blending mandates as well as specific mentions of ethanol (or other biofuels) are covered here.",
    "A_Finance": "This parameter records financing instruments used to pay for technologies, projects and programmes that reduce GHG emissions, e.g. climate finance solutions, investments in EVs, green bonds. Not to be confused with economic instruments.",
    "R_System": "This parameter identifies efforts to adapt the transport system to climate-change impacts and to increase its resilience.",
    "R_Infrares": "This parameter identifies efforts to adapt transport infrastructure to climate-change impacts and to increase its resilience.",
}

# Source-string corrections (typos in the Source column itself)
source_corrections = {
    "VTI (215): Comprehensive Transport Planning": "VTI (2015): Comprehensive Transport Planning",
}

for rec in records:
    key = rec["code"] or rec["label"]
    sug = suggestions.get(key)
    rec["definition_suggested"] = sug
    if rec["source"] in source_corrections:
        rec["source_original"] = rec["source"]
        rec["source"] = source_corrections[rec["source"]]


# ---------------------------------------------------------------------------
# 4. Cross-cutting dimensions taken from the live database column structure
#    (Mode, Geography, Activity, Status). These are tagging dimensions applied
#    to mitigation/adaptation rows, captured here for completeness.
# ---------------------------------------------------------------------------
cross_cutting = {
    "mode": {
        "description": "Transport mode(s) a measure applies to (multi-select; columns flagged 'x' in the database).",
        "applies_to": ["Mitigation", "Adaptation"],
        "values": [
            {"label": "Mode: Not defined"},
            {"label": "Informal transport"},
            {"label": "Active mobility", "children": ["Walking", "Cycling"]},
            {"label": "Road", "children": ["Two-/Three-wheelers", "Cars", "Private cars", "Taxis", "Truck", "Bus"]},
            {"label": "Rail", "children": ["Heavy rail", "High-speed rail", "Transit rail"]},
            {"label": "Water", "children": ["Coastal shipping", "Inland shipping", "International maritime"]},
            {"label": "Aviation", "children": ["Domestic aviation", "International aviation"]},
        ],
    },
    "geography": {
        "description": "Spatial setting a measure applies to (multi-select).",
        "applies_to": ["Mitigation", "Adaptation"],
        "values": [
            {"label": "Geography: Not defined"},
            {"label": "Urban"},
            {"label": "Rural"},
            {"label": "Inter-city"},
        ],
    },
    "activity": {
        "description": "Whether a measure targets passenger or freight transport.",
        "applies_to": ["Mitigation", "Adaptation"],
        "values": [
            {"label": "Not defined", "code": "E_Notdefined"},
            {"label": "Passenger transport", "code": "E_Passenger"},
            {"label": "Freight", "code": "E_Freight"},
            {"label": "Both passenger and freight", "code": "E_Passenger&Freight"},
        ],
    },
    "status_of_measure": {
        "description": "Implementation status of a mitigation measure.",
        "applies_to": ["Mitigation"],
        "values": [
            {"label": "Implemented or adopted", "code": "S_Implemented"},
            {"label": "Planned", "code": "S_Planned"},
            {"label": "Unclear", "code": "S_Unclear"},
        ],
    },
    "asi": {
        "description": "Avoid-Shift-Improve framework classification of a mitigation measure.",
        "applies_to": ["Mitigation"],
        "values": [
            {"label": "Avoid"}, {"label": "Shift"}, {"label": "Improve"},
            {"label": "Avoid, shift"}, {"label": "Avoid, improve"},
            {"label": "Shift, improve"}, {"label": "Avoid, shift, improve"},
        ],
    },
}

# Pull ASI / Activity / Status definitions out of records into cross_cutting,
# so the main records list holds only the primary classification tree.
def pull(domain, dimension):
    kept, moved = [], []
    for rec in records:
        if rec["domain"] == domain and rec["dimension"] == dimension:
            moved.append(rec)
        else:
            kept.append(rec)
    return kept, moved

asi_defs = {r["label"]: r["definition"] for r in records if r["dimension"] == "A-S-I"}
for v in cross_cutting["asi"]["values"]:
    v["definition"] = asi_defs.get(v["label"])

act_defs = {r["label"]: r["definition"] for r in records if r["dimension"] == "Activity"}
for v in cross_cutting["activity"]["values"]:
    v["definition"] = act_defs.get(v["label"])

stat_defs = {r["label"]: r["definition"] for r in records if r["dimension"] == "Status of measure"}
for v in cross_cutting["status_of_measure"]["values"]:
    v["definition"] = stat_defs.get(v["label"])

# Remove the pulled dimensions from the primary records
records = [r for r in records if r["dimension"] not in ("A-S-I", "Activity", "Status of measure")]


# ---------------------------------------------------------------------------
# 5. Assemble final JSON object
# ---------------------------------------------------------------------------
taxonomy = {
    "name": "NDC Transport Tracker Taxonomy",
    "version": "4.0",
    "source_workbook": "NDC-taxonomy.xlsx",
    "source_update_date": "2026-05-05",
    "maintainers": [
        {"name": "Maria Belen Vasquez", "org": "GIZ", "email": "maria.vasquez1@giz.de"},
        {"name": "Nikola Medimorec", "org": "SLOCAT", "email": "nikola.medimorec@slocatpartnership.org"},
    ],
    "license": "CC BY 4.0",
    "citation": "GIZ and SLOCAT (2025). NDC Transport Tracker (vers. 4.0). Available from: www.changing-transport.org/tracker.",
    "definition_policy": "Definitions are preserved verbatim in 'definition'. Where a clarified version is offered it appears in 'definition_suggested' (null when no change is suggested). The maintainer decides which to adopt. Source typos are preserved in 'source_original' when corrected.",
    "domains": {},
    "cross_cutting_dimensions": cross_cutting,
}

for rec in records:
    d = rec.pop("domain")
    dim = rec["dimension"]
    taxonomy["domains"].setdefault(d, {})
    taxonomy["domains"][d].setdefault(dim, [])
    taxonomy["domains"][d][dim].append(rec)

with open(OUT / "ndc_taxonomy.json", "w", encoding="utf-8") as f:
    json.dump(taxonomy, f, ensure_ascii=False, indent=2)
print("Wrote ndc_taxonomy.json  —", sum(len(v) for d in taxonomy['domains'].values() for v in d.values()), "domain records")


# ===========================================================================
# 6. Generate flat CSV and human-readable README (Mermaid) from the same data
# ===========================================================================
T = taxonomy

# ---------------------------------------------------------------------------
# CSV: one row per classification value, flat.
# ---------------------------------------------------------------------------
fields = ["domain", "dimension", "label", "code", "parent_1", "parent_2",
          "definition", "definition_suggested", "source"]
with open(str(OUT / "ndc_taxonomy.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
    w.writeheader()
    for domain, dims in T["domains"].items():
        for dim, recs in dims.items():
            for r in recs:
                row = dict(r)
                row["domain"] = domain
                w.writerow(row)
    # cross-cutting
    for cc_key, cc in T["cross_cutting_dimensions"].items():
        for v in cc["values"]:
            w.writerow({
                "domain": "Cross-cutting",
                "dimension": cc_key,
                "label": v["label"],
                "code": v.get("code"),
                "definition": v.get("definition"),
            })
            for child in v.get("children", []):
                w.writerow({
                    "domain": "Cross-cutting",
                    "dimension": cc_key,
                    "label": child,
                    "parent_1": v["label"],
                })

print("Wrote ndc_taxonomy.csv")

# ---------------------------------------------------------------------------
# Mermaid helpers
# ---------------------------------------------------------------------------
def mid(s):
    """Make a safe mermaid node id."""
    return "n" + str(abs(hash(s)) % (10**9))

def esc(s):
    return s.replace('"', "'").replace("\n", " ").strip()

def code_to_label():
    m = {}
    for domain, dims in T["domains"].items():
        for dim, recs in dims.items():
            for r in recs:
                if r.get("code"):
                    m[r["code"]] = r["label"]
    return m

C2L = code_to_label()

# Mitigation tree: Category (code C_*) -> Purpose (P_*, parent_1=category) ->
# Instrument (parent_1=category, parent_2=purpose)
def mermaid_mitigation():
    cats = T["domains"]["Mitigation"]["Type of measure: category"]
    purps = T["domains"]["Mitigation"]["Type of measure: purpose"]
    insts = T["domains"]["Mitigation"]["Type of measure: instruments"]
    lines = ["```mermaid", "graph LR"]
    for c in cats:
        cn = mid(c["code"])
        lines.append(f'  {cn}["{esc(c["label"])}"]')
        for p in purps:
            if p["parent_1"] == c["code"]:
                pn = mid(p["code"])
                lines.append(f'  {pn}("{esc(p["label"])}")')
                lines.append(f"  {cn} --> {pn}")
                for ins in insts:
                    if ins["parent_2"] == p["code"]:
                        inn = mid(ins["code"] + p["code"])
                        lines.append(f'  {pn} --> {inn}["{esc(ins["label"])}"]')
    lines.append("```")
    return "\n".join(lines)

def mermaid_targets():
    lines = ["```mermaid", "graph LR", '  T["Targets"]']
    dims = ["Target area", "Conditionality", "GHG target?", "Target type",
            "Target scope", "ICE phase-out target type"]
    for dim in dims:
        dn = mid(dim)
        lines.append(f'  T --> {dn}("{esc(dim)}")')
        for r in T["domains"]["Targets"][dim]:
            rn = mid(r["label"] + dim)
            lines.append(f'  {dn} --> {rn}["{esc(r["label"])}"]')
    lines.append("```")
    return "\n".join(lines)

def mermaid_adaptation():
    lines = ["```mermaid", "graph LR"]
    cats = T["domains"]["Adaptation"]["Category"]
    meas = T["domains"]["Adaptation"]["Measure"]
    for c in cats:
        cn = mid(c["code"])
        lines.append(f'  {cn}["{esc(c["label"])}"]')
        for m in meas:
            if m["parent_1"] == c["code"]:
                mn = mid(m["code"])
                lines.append(f'  {cn} --> {mn}["{esc(m["label"])}"]')
    lines.append("```")
    return "\n".join(lines)

def mermaid_benefits():
    lines = ["```mermaid", "graph LR", '  B["Benefits"]']
    for r in T["domains"]["Benefits"]["Type of benefit"]:
        rn = mid(r["label"])
        lines.append(f'  B --> {rn}["{esc(r["label"])}"]')
    lines.append("```")
    return "\n".join(lines)

# ---------------------------------------------------------------------------
# Definition tables (collapsible per dimension)
# ---------------------------------------------------------------------------
def def_table(recs, show_parents=True):
    out = []
    out.append("| Label | Code | " + ("Parent | " if show_parents else "") + "Definition | Suggested clarification | Source |")
    out.append("|---|---|" + ("---|" if show_parents else "") + "---|---|---|")
    for r in recs:
        parent = " / ".join([p for p in [r.get("parent_1"), r.get("parent_2")] if p]) or ""
        d = (r.get("definition") or "").replace("|", "\\|").replace("\n", " ")
        s = (r.get("definition_suggested") or "").replace("|", "\\|").replace("\n", " ")
        src = (r.get("source") or "").replace("|", "\\|")
        code = r.get("code") or ""
        row = f"| {r['label']} | `{code}` | "
        if show_parents:
            row += f"{parent} | "
        row += f"{d} | {s} | {src} |"
        out.append(row)
    return "\n".join(out)

# ---------------------------------------------------------------------------
# Assemble README
# ---------------------------------------------------------------------------
md = []
md.append("# NDC Transport Tracker — Taxonomy\n")
md.append(f"**Version {T['version']}** · source workbook `{T['source_workbook']}` · last updated {T['source_update_date']} · licensed {T['license']}\n")
md.append("> " + T["citation"] + "\n")
md.append("This document is the full classification taxonomy used by the NDC Transport Tracker and the Transport Policy Miner pipeline. It is available in three forms in this directory:\n")
md.append("- [`ndc_taxonomy.json`](./ndc_taxonomy.json) — canonical, machine-readable (feed this to code or an LLM)\n- [`ndc_taxonomy.csv`](./ndc_taxonomy.csv) — flat, spreadsheet-friendly\n- this file — human-readable, with rendered diagrams\n")

md.append("## Definition policy\n")
md.append(T["definition_policy"] + "\n")

md.append("## Structure at a glance\n")
md.append("The taxonomy has four **domains** (Targets, Mitigation, Adaptation, Benefits). Each domain is split into one or more **dimensions**. Some dimensions are flat lists of values; others are hierarchical (a value opens into child values). On top of these, several **cross-cutting dimensions** (mode, geography, activity, status, Avoid-Shift-Improve) are tagged onto individual measures.\n")

md.append("```mermaid\ngraph TD\n  R[NDC Transport Tracker Taxonomy] --> T[Targets]\n  R --> M[Mitigation]\n  R --> A[Adaptation]\n  R --> B[Benefits]\n  R --> X[Cross-cutting dimensions]\n  M --> MC[Category] --> MP[Purpose] --> MI[Instrument]\n  A --> AC[Category] --> AM[Measure]\n  X --> mode & geography & activity & status & ASI[Avoid-Shift-Improve]\n```\n")

# Targets
md.append("---\n## 1. Targets\n")
md.append("A target expresses a goal or objective related to lowering GHG emissions or adapting to climate impacts. Targets are classified across six independent dimensions.\n")
md.append(mermaid_targets() + "\n")
for dim in ["Target area", "Conditionality", "GHG target?", "Target type", "Target scope", "ICE phase-out target type"]:
    recs = T["domains"]["Targets"][dim]
    md.append(f"<details>\n<summary><b>{dim}</b> ({len(recs)} values)</summary>\n")
    md.append(def_table(recs) + "\n")
    md.append("</details>\n")

# Mitigation
md.append("---\n## 2. Mitigation measures\n")
md.append("A mitigation measure is an action or pathway that leads to lowering GHG emissions. Mitigation is a three-level hierarchy: **Category → Purpose → Instrument**. An instrument always belongs to one purpose, and a purpose to one category.\n")
md.append(mermaid_mitigation() + "\n")
md.append("<details>\n<summary><b>Categories</b> (top level)</summary>\n")
md.append(def_table(T["domains"]["Mitigation"]["Type of measure: category"], show_parents=False) + "\n</details>\n")
md.append("<details>\n<summary><b>Purposes</b> (mid level, with parent category)</summary>\n")
md.append(def_table(T["domains"]["Mitigation"]["Type of measure: purpose"]) + "\n</details>\n")
md.append("<details>\n<summary><b>Instruments</b> (leaf level, with parent category / purpose)</summary>\n")
md.append(def_table(T["domains"]["Mitigation"]["Type of measure: instruments"]) + "\n</details>\n")

# Adaptation
md.append("---\n## 3. Adaptation measures\n")
md.append("An adaptation measure is a policy or strategy that reduces the risks and harms caused by climate change. Adaptation is a two-level hierarchy: **Category → Measure**.\n")
md.append(mermaid_adaptation() + "\n")
md.append("<details>\n<summary><b>Categories</b></summary>\n")
md.append(def_table(T["domains"]["Adaptation"]["Category"], show_parents=False) + "\n</details>\n")
md.append("<details>\n<summary><b>Measures</b> (with parent category)</summary>\n")
md.append(def_table(T["domains"]["Adaptation"]["Measure"]) + "\n</details>\n")

# Benefits
md.append("---\n## 4. Benefits\n")
md.append("A benefit links a climate-related transport action to other positive impacts. Benefits are a single flat dimension (tagged per measure).\n")
md.append(mermaid_benefits() + "\n")
md.append("<details>\n<summary><b>Type of benefit</b></summary>\n")
md.append(def_table(T["domains"]["Benefits"]["Type of benefit"], show_parents=False) + "\n</details>\n")

# Cross-cutting
md.append("---\n## 5. Cross-cutting dimensions\n")
md.append("These dimensions are tagged onto individual mitigation and/or adaptation rows in the database (multi-select where noted). They are not part of the primary classification tree but are needed to reproduce the database schema.\n")
for cc_key, cc in T["cross_cutting_dimensions"].items():
    md.append(f"<details>\n<summary><b>{cc_key}</b> — {esc(cc['description'])}</summary>\n")
    md.append(f"*Applies to: {', '.join(cc['applies_to'])}*\n")
    md.append("| Value | Code | Sub-values | Definition |")
    md.append("|---|---|---|---|")
    for v in cc["values"]:
        children = ", ".join(v.get("children", []))
        d = (v.get("definition") or "").replace("|", "\\|")
        code = v.get("code") or ""
        md.append(f"| {v['label']} | `{code}` | {children} | {d} |")
    md.append("\n</details>\n")

with open(str(OUT / "TAXONOMY.md"), "w", encoding="utf-8") as f:
    f.write("\n".join(md))

print("Wrote TAXONOMY.md")

print("\nBuild complete. Outputs written to", OUT)
